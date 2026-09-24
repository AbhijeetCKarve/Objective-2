"""Shared code for the corrected MAML-AE notebooks (Notebook_01 to Notebook_10).

Every notebook imports from this file instead of re-defining the data handling, the
model, the meta-learning loop or the metrics. The original notebooks each carried their own
copy of this code, which is how a broken training loop (Kaggle notebook
"04-maml-training (3)") went unnoticed.
"""
import ast
import copy
import json
import os
import subprocess
import time

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
from sklearn.metrics import (average_precision_score, f1_score, precision_recall_curve,
                             precision_score, recall_score, roc_auc_score)
from sklearn.preprocessing import MinMaxScaler

WINDOW = 30
STRIDE_NORMAL = 10     # normal windows from the training file (as in the original 02)
STRIDE_ANOMALY = 5     # legacy anomaly windows from the test file (as in the original 02)

# ---------------------------------------------------------------------------
# Channel split. This is the exact split the original work used
# (task_splits_25feat.json). It is written out here so it no longer depends on the
# order in which cells of the original 02/04 were run.
# ---------------------------------------------------------------------------
META_TRAIN = ["P-3", "A-9", "G-2", "G-7", "E-5", "S-1", "T-3", "G-4", "E-10", "B-1",
              "D-13", "E-1", "A-7", "A-2", "T-1", "G-6", "A-4", "D-9", "D-4", "P-1",
              "P-2", "D-2", "F-1", "E-11", "P-7", "E-9", "E-2", "E-4", "G-3", "F-3",
              "E-13", "A-3", "R-1", "F-2", "D-8", "A-5", "A-1", "D-1", "E-8"]
META_VAL = ["G-1", "E-12", "D-5", "E-7", "A-8", "D-11"]
META_TEST = ["E-3", "D-7", "E-6", "D-12", "D-6", "T-2", "A-6", "D-3", "P-4"]
EXCLUDED = {
    "D-12": "only 312 training timesteps (29 normal windows at stride 10): too few to "
            "draw support sets and hold out normal data. Label-free reason.",
    "P-4": "sensor-dropout (flat-line) anomaly that a reconstruction model cannot flag. "
           "This reason was found by inspecting the test anomalies in the original "
           "notebook 04, so the exclusion is post hoc. P-4 is still scored and reported "
           "separately as a sensitivity row.",
}
EVAL_CHANNELS = [c for c in META_TEST if c not in EXCLUDED]
SENSITIVITY_CHANNELS = ["P-4"]


# ---------------------------------------------------------------------------
# Paths, seeds, provenance
# ---------------------------------------------------------------------------
def find_file(name, extra_roots=()):
    """Search /kaggle/input, then the given roots, for a file called `name`."""
    roots = [r for r in ["/kaggle/input", *extra_roots] if r and os.path.isdir(r)]
    for root in roots:
        for dirpath, _, files in os.walk(root):
            if name in files:
                return os.path.join(dirpath, name)
    return None


def data_roots(extra_roots=()):
    """Where raw data is searched for: /kaggle/input (always), the folders listed in the
    MAML_DATA_ROOT environment variable (separated by os.pathsep), SMAP_ROOT, and the
    folder above this repository (so the data folders of a neighbouring checkout work)."""
    here = os.path.dirname(os.path.abspath(__file__))
    env = [r for r in os.environ.get("MAML_DATA_ROOT", "").split(os.pathsep) if r]
    return [os.environ.get("SMAP_ROOT"), *extra_roots, *env, os.path.dirname(here)]


def find_data_file(name, extra_roots=()):
    return find_file(name, data_roots(extra_roots))


def find_smap_raw(extra_roots=()):
    """Locate labeled_anomalies.csv and the train/ and test/ folders of the SMAP release."""
    roots = data_roots(extra_roots)
    csv = find_file("labeled_anomalies.csv", roots)
    if csv is None:
        raise FileNotFoundError("labeled_anomalies.csv not found under /kaggle/input or "
                                f"{roots}. Set MAML_DATA_ROOT to the folder holding the data.")
    base = os.path.dirname(csv)
    for dirpath, dirs, _ in os.walk(base):
        if "train" in dirs and "test" in dirs:
            return csv, os.path.join(dirpath, "train"), os.path.join(dirpath, "test")
    raise FileNotFoundError(f"train/ and test/ folders not found below {base}")


def output_dir(smoke=False):
    """/kaggle/working on Kaggle, otherwise a local folder next to this file. Smoke tests
    write to a separate SMOKE subfolder so their files are never mistaken for results."""
    if os.path.isdir("/kaggle/working"):
        d = "/kaggle/working"
    else:
        d = os.path.join(os.path.dirname(os.path.abspath(__file__)), "outputs")
    if smoke:
        d = os.path.join(d, "SMOKE")
    os.makedirs(d, exist_ok=True)
    return d


def seed_everything(seed):
    np.random.seed(seed)
    torch.manual_seed(seed)


def git_commit():
    """Commit hash of this code. Kaggle has no git, so a CODE_VERSION file is used there."""
    here = os.path.dirname(os.path.abspath(__file__))
    try:
        return subprocess.check_output(["git", "-C", here, "rev-parse", "HEAD"],
                                       stderr=subprocess.DEVNULL).decode().strip()
    except Exception:
        p = find_file("CODE_VERSION", [here])
        return open(p).read().strip() if p else "unknown"


def save_json(path, obj):
    obj = dict(obj)
    obj.setdefault("git_commit", git_commit())
    obj.setdefault("timestamp", time.strftime("%Y-%m-%dT%H:%M:%S"))
    with open(path, "w") as f:
        json.dump(obj, f, indent=2, default=_json_default)
    return path


def _json_default(o):
    if isinstance(o, (np.integer,)):
        return int(o)
    if isinstance(o, (np.floating,)):
        return float(o)
    if isinstance(o, np.ndarray):
        return o.tolist()
    raise TypeError(type(o))


# ---------------------------------------------------------------------------
# SMAP labels and data
# ---------------------------------------------------------------------------
def load_labels(csv_path):
    """One row per channel. The CSV lists P-2 twice (two different anomaly ranges);
    the rows are merged here. The original code kept only the first P-2 row."""
    df = pd.read_csv(csv_path)
    df["seqs"] = df["anomaly_sequences"].apply(ast.literal_eval)   # no eval()
    rows = []
    for ch, g in df.groupby("chan_id", sort=False):
        seqs = [list(map(int, s)) for ss in g["seqs"] for s in ss]   # CSV order kept
        rows.append({"chan_id": ch, "spacecraft": g["spacecraft"].iloc[0],
                     "sequences": seqs, "n_csv_rows": len(g),
                     "num_values": int(g["num_values"].iloc[0])})
    return pd.DataFrame(rows).set_index("chan_id")


def point_labels(sequences, n):
    """Per-timestep labels. Sequence ends in labeled_anomalies.csv are inclusive."""
    y = np.zeros(n, dtype=np.int8)
    for s, e in sequences:
        y[max(0, s):min(n, e + 1)] = 1
    return y


def create_windows(data, window=WINDOW, stride=1):
    n = (len(data) - window) // stride + 1
    if n <= 0:
        return np.zeros((0, window, data.shape[1]), dtype=np.float32)
    idx = np.arange(window)[None, :] + stride * np.arange(n)[:, None]
    return data[idx].astype(np.float32)


def prepare_channel(train_raw, test_raw, sequences, clip=True):
    """MinMax scaling fitted on the training file only, then windows and labels."""
    scaler = MinMaxScaler().fit(train_raw)
    tr = scaler.transform(train_raw)
    te = scaler.transform(test_raw)
    if clip:
        tr, te = np.clip(tr, 0, 1), np.clip(te, 0, 1)
    tr, te = tr.astype(np.float32), te.astype(np.float32)
    y = point_labels(sequences, len(te))
    normal_windows = create_windows(tr, WINDOW, STRIDE_NORMAL)
    normal_starts = np.arange(len(normal_windows)) * STRIDE_NORMAL
    return {"train": tr, "test": te, "labels": y,
            "normal_windows": normal_windows, "normal_starts": normal_starts,
            "legacy_anomaly_windows": legacy_anomaly_windows(te, sequences),
            "scaler_min": scaler.data_min_.astype(np.float32),
            "scaler_max": scaler.data_max_.astype(np.float32)}


def build_smap_dataset(channels=None, extra_roots=()):
    """Load and prepare SMAP channels straight from the raw release (a few seconds).
    Returns ({channel: prepared dict}, labels table). Only SMAP channels (25 features)."""
    csv, train_dir, test_dir = find_smap_raw(extra_roots)
    labels = load_labels(csv)
    if channels is None:
        channels = [c for c in labels.index if labels.loc[c, "spacecraft"] == "SMAP"]
    data = {}
    for ch in channels:
        data[ch] = prepare_channel(np.load(os.path.join(train_dir, f"{ch}.npy")),
                                   np.load(os.path.join(test_dir, f"{ch}.npy")),
                                   labels.loc[ch, "sequences"])
    return data, labels


def legacy_anomaly_windows(test_scaled, sequences):
    """Anomaly windows exactly as the original 02 built them: for each labelled segment,
    stride-5 windows over test[start : end + 30]."""
    out = []
    for s, e in sequences:
        seg = test_scaled[max(0, s):min(e + WINDOW, len(test_scaled))]
        if len(seg) >= WINDOW:
            out.append(create_windows(seg, WINDOW, STRIDE_ANOMALY))
    return np.concatenate(out) if out else np.zeros((0, WINDOW, test_scaled.shape[1]), np.float32)


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
class LSTMEncoder(nn.Module):
    def __init__(self, input_size=25, hidden1=64, hidden2=32, latent=16):
        super().__init__()
        self.lstm1 = nn.LSTM(input_size, hidden1, batch_first=True)
        self.lstm2 = nn.LSTM(hidden1, hidden2, batch_first=True)
        self.fc = nn.Linear(hidden2, latent)

    def forward(self, x):
        out, _ = self.lstm1(x)
        _, (h_n, _) = self.lstm2(out)       # keep only the last hidden state
        return self.fc(h_n[-1])


class LSTMDecoder(nn.Module):
    def __init__(self, latent=16, hidden1=32, hidden2=64, output_size=25, seq_len=WINDOW):
        super().__init__()
        self.seq_len = seq_len
        self.lstm1 = nn.LSTM(latent, hidden1, batch_first=True)
        self.lstm2 = nn.LSTM(hidden1, hidden2, batch_first=True)
        self.fc = nn.Linear(hidden2, output_size)

    def forward(self, z):
        z_rep = z.unsqueeze(1).repeat(1, self.seq_len, 1)
        out, _ = self.lstm1(z_rep)
        out, _ = self.lstm2(out)
        return self.fc(out)


class LSTMAutoencoder(nn.Module):
    """Encoder 64 -> 32 -> code 16; decoder 16 -> 32 -> 64 -> d. 69,481 parameters at d=25."""

    def __init__(self, input_size=25, seq_len=WINDOW, hidden1=64, hidden2=32, latent=16):
        super().__init__()
        self.encoder = LSTMEncoder(input_size, hidden1, hidden2, latent)
        self.decoder = LSTMDecoder(latent, hidden2, hidden1, input_size, seq_len)

    def forward(self, x):
        return self.decoder(self.encoder(x))

    def reconstruction_error(self, x):
        return torch.mean((x - self.forward(x)) ** 2, dim=(1, 2))


class MLPAutoencoder(nn.Module):
    def __init__(self, input_size=25, seq_len=WINDOW, latent=16):
        super().__init__()
        self.seq_len, self.input_size = seq_len, input_size
        flat = seq_len * input_size
        self.encoder = nn.Sequential(nn.Linear(flat, 256), nn.ReLU(), nn.Linear(256, 64),
                                     nn.ReLU(), nn.Linear(64, latent))
        self.decoder = nn.Sequential(nn.Linear(latent, 64), nn.ReLU(), nn.Linear(64, 256),
                                     nn.ReLU(), nn.Linear(256, flat))

    def forward(self, x):
        b = x.shape[0]
        return self.decoder(self.encoder(x.reshape(b, -1))).reshape(b, self.seq_len, self.input_size)

    def reconstruction_error(self, x):
        return torch.mean((x - self.forward(x)) ** 2, dim=(1, 2))


def count_params(model):
    return sum(p.numel() for p in model.parameters() if p.requires_grad)


# ---------------------------------------------------------------------------
# First-order MAML (FOMAML)
# ---------------------------------------------------------------------------
_mse = nn.MSELoss()


def inner_adapt(model, support, lr=0.01, steps=10):
    """Deep-copy the model and take `steps` SGD steps on the support windows.
    steps=0 returns an unchanged copy (the zero-step control)."""
    learner = copy.deepcopy(model)
    learner.train()
    opt = torch.optim.SGD(learner.parameters(), lr=lr)
    for _ in range(steps):
        opt.zero_grad()
        _mse(learner(support), support).backward()
        opt.step()
    return learner


def fomaml_outer_step(model, outer_opt, episodes, inner_lr=0.01, inner_steps=10,
                      clip=1.0, train=True):
    """One outer step. `episodes` is a list of (support, query) tensors.

    The query-loss gradient is taken with respect to the ADAPTED copy's parameters,
    averaged over tasks, and written onto the meta-parameters' .grad before the
    optimizer step. Calling .backward() on the copy instead leaves the meta-model
    without gradients; check_outer_step() catches that."""
    params = list(model.parameters())
    accum = [torch.zeros_like(p) for p in params]
    total = 0.0
    for support, query in episodes:
        learner = inner_adapt(model, support, inner_lr, inner_steps)
        qloss = _mse(learner(query), query)
        if train:
            grads = torch.autograd.grad(qloss, list(learner.parameters()))
            for a, g in zip(accum, grads):
                a.add_(g.detach())
        total += qloss.item()
    if train:
        outer_opt.zero_grad()
        for p, a in zip(params, accum):
            p.grad = a / len(episodes)
        torch.nn.utils.clip_grad_norm_(params, clip)
        outer_opt.step()
    return total / len(episodes)


def check_outer_step(model, before):
    """Runtime checks 1 and 2: every meta-parameter has a non-zero gradient, and the
    meta-parameters changed. `before` is a list of parameter copies taken before the step."""
    grads = [p.grad for p in model.parameters()]
    missing = [i for i, g in enumerate(grads) if g is None or float(g.abs().sum()) == 0.0]
    moved = sum(float((p.detach() - b).norm() ** 2) for p, b in zip(model.parameters(), before)) ** 0.5
    assert not missing, f"meta-parameters without gradient: {missing}"
    assert moved > 0, "meta-parameters did not change after the outer step"
    return {"n_params_with_grad": len(grads) - len(missing), "meta_param_change_l2": moved}


def adaptation_report(model, adapted, support):
    """Check 4: how far adaptation moved the parameters, and whether support loss fell."""
    num = sum(float((a.detach() - p.detach()).norm() ** 2)
              for a, p in zip(adapted.parameters(), model.parameters())) ** 0.5
    den = sum(float(p.detach().norm() ** 2) for p in model.parameters()) ** 0.5
    with torch.no_grad():
        before = _mse(model(support), support).item()
        after = _mse(adapted(support), support).item()
    return {"relative_param_change": num / den, "support_loss_before": before,
            "support_loss_after": after}


def sample_episode(windows, rng, support_size=20, query_size=20):
    idx = rng.permutation(len(windows))
    need = support_size + query_size
    if len(windows) < need:
        s = windows[rng.choice(len(windows), support_size, replace=True)]
        q = windows[rng.choice(len(windows), query_size, replace=True)]
    else:
        s, q = windows[idx[:support_size]], windows[idx[support_size:need]]
    return s, q


def fixed_episodes(windows_by_task, n_per_task, seed, support_size=20, query_size=20):
    """Validation episodes drawn once with their own seed and reused at every check.
    The original drew one fresh random episode per check from the training RNG, which
    made the validation curve noisy and coupled it to the training sample stream."""
    rng = np.random.RandomState(seed)
    return [(t, *sample_episode(w, rng, support_size, query_size))
            for t, w in windows_by_task.items() for _ in range(n_per_task)]


def to_tensor(x, device):
    return torch.as_tensor(np.asarray(x), dtype=torch.float32, device=device)


def meta_validation_loss(model, val_episodes, device, inner_lr, inner_steps):
    eps = [(to_tensor(s, device), to_tensor(q, device)) for _, s, q in val_episodes]
    return fomaml_outer_step(model, None, eps, inner_lr, inner_steps, train=False)


def train_maml(model, train_windows, val_episodes, device, *, seed, n_outer, val_every,
               inner_lr=0.01, inner_steps=10, outer_lr=1e-3, tasks_per_batch=4,
               support_size=20, query_size=20, use_scheduler=True, ckpt_path=None,
               resume_from=None, patience=None, log=print):
    """Meta-train with FOMAML. Training episodes come from their own RNG (seeded by
    `seed`); validation uses the fixed `val_episodes`, so validating does not change the
    training sample stream. The best-validation state is returned. Checkpoints hold the
    RNG state too, so a resumed run continues the same sample stream."""
    tasks = list(train_windows)
    rng = np.random.RandomState(seed)
    opt = torch.optim.Adam(model.parameters(), lr=outer_lr)
    sched = (torch.optim.lr_scheduler.ReduceLROnPlateau(opt, "min", factor=0.5, patience=5,
                                                        min_lr=1e-5) if use_scheduler else None)
    state = {"step": 0, "best_val": float("inf"), "best_step": 0, "history": [],
             "best_state": copy.deepcopy(model.state_dict()), "checks": {}, "bad_checks": 0,
             "stopped_early": False}
    if resume_from is not None:
        ck = torch.load(resume_from, map_location=device, weights_only=False)
        model.load_state_dict(ck["model"]); opt.load_state_dict(ck["opt"])
        if sched is not None and ck.get("sched") is not None:
            sched.load_state_dict(ck["sched"])
        rng.set_state(ck["rng"])
        state.update({k: ck[k] for k in ["step", "best_val", "best_step", "history",
                                         "best_state", "checks", "bad_checks",
                                         "stopped_early"] if k in ck})
        log(f"resumed from {resume_from} at step {state['step']}")
    t0 = time.time()
    first_val = None
    for step in range(state["step"] + 1, n_outer + 1):
        if state["stopped_early"]:
            break
        batch = rng.choice(tasks, size=min(tasks_per_batch, len(tasks)), replace=False)
        eps = []
        for t in batch:
            s, q = sample_episode(train_windows[t], rng, support_size, query_size)
            eps.append((to_tensor(s, device), to_tensor(q, device)))
        before = [p.detach().clone() for p in model.parameters()] if step == 1 else None
        model.train()
        tl = fomaml_outer_step(model, opt, eps, inner_lr, inner_steps)
        if step == 1:        # runtime checks 1, 2 and 4
            state["checks"]["outer_step"] = check_outer_step(model, before)
            s, _ = eps[0]
            rep = adaptation_report(model, inner_adapt(model, s, inner_lr, inner_steps), s)
            state["checks"]["adaptation_at_step_1"] = rep
            assert rep["relative_param_change"] > 0, "inner loop did not move the parameters"
            if rep["support_loss_after"] >= rep["support_loss_before"]:
                log("WARNING: adaptation did not lower the support loss at step 1")
        if step % val_every == 0 or step == n_outer:
            vl = meta_validation_loss(model, val_episodes, device, inner_lr, inner_steps)
            if sched is not None:
                sched.step(vl)
            state["history"].append({"step": step, "train_loss": tl, "val_loss": vl,
                                     "lr": opt.param_groups[0]["lr"], "seconds": time.time() - t0})
            first_val = state["history"][0]["val_loss"]
            if vl < state["best_val"] - 1e-6:
                state.update(best_val=vl, best_step=step, bad_checks=0,
                             best_state=copy.deepcopy(model.state_dict()))
            else:
                state["bad_checks"] += 1
                if patience is not None and state["bad_checks"] >= patience:
                    state["stopped_early"] = True
                    log(f"early stop at step {step}: no improvement in {patience} checks")
            log(f"step {step:6d} | train {tl:.6f} | val {vl:.6f} | best {state['best_val']:.6f} "
                f"@ {state['best_step']} | {time.time() - t0:.0f}s")
            state["step"] = step
            if ckpt_path:
                torch.save({"model": model.state_dict(), "opt": opt.state_dict(),
                            "sched": sched.state_dict() if sched else None,
                            "rng": rng.get_state(), **state}, ckpt_path)
    vals = [h["val_loss"] for h in state["history"]]
    flat = len(vals) > 1 and max(vals) - min(vals) <= 1e-9 * max(1.0, abs(first_val or 0))
    if flat:   # runtime check 3
        log("WARNING: meta-validation loss is flat. The meta-model may not be training.")
    state["checks"]["val_loss_flat"] = bool(flat)
    model.load_state_dict(state["best_state"])
    info = {k: state[k] for k in ["best_val", "best_step", "history", "checks", "stopped_early"]}
    info.update(steps_reached=state["step"], max_outer_steps=n_outer, patience=patience)
    return model, info


def train_conventional(model, pool, device, *, seed, max_epochs=150, patience=15, batch=128,
                       lr=1e-3, log=print):
    """Conventional training on pooled normal windows (the original static-AE recipe:
    random 90/10 split, Adam 1e-3, batch 128, early stopping on the 10% split)."""
    seed_everything(seed)
    p = np.random.RandomState(seed).permutation(len(pool))
    cut = int(0.9 * len(pool))
    tr = torch.as_tensor(pool[p[:cut]], dtype=torch.float32)
    va = to_tensor(pool[p[cut:]], device)
    opt = torch.optim.Adam(model.parameters(), lr=lr)
    g = torch.Generator().manual_seed(seed)
    best, bad, best_state, hist = float("inf"), 0, None, []
    for ep in range(max_epochs):
        model.train()
        perm = torch.randperm(len(tr), generator=g)
        for st in range(0, len(tr), batch):
            b = tr[perm[st:st + batch]].to(device)
            opt.zero_grad(); _mse(model(b), b).backward(); opt.step()
        model.eval()
        with torch.no_grad():
            vl = _mse(model(va), va).item()
        hist.append(vl)
        if vl < best - 1e-6:
            best, bad, best_state = vl, 0, copy.deepcopy(model.state_dict())
        else:
            bad += 1
            if bad >= patience:
                break
    model.load_state_dict(best_state)
    log(f"conventional training: best val {best:.6f} after {len(hist)} epochs")
    return model, {"best_val": best, "epochs_run": len(hist), "val_history": hist}


# ---------------------------------------------------------------------------
# Scoring and metrics
# ---------------------------------------------------------------------------
@torch.no_grad()
def window_errors(model, windows, device, batch=2048):
    model.eval()
    out = []
    for i in range(0, len(windows), batch):
        out.append(model.reconstruction_error(to_tensor(windows[i:i + batch], device)).cpu().numpy())
    return np.concatenate(out) if out else np.zeros(0, np.float32)


def pointwise_from_window_errors(errors, n, window=WINDOW):
    """Stride-1 windows w = 0..n-window. A timestep's score is the mean error of all
    windows that cover it (between 1 and `window` windows)."""
    assert len(errors) == n - window + 1
    diff = np.zeros(n + 1)
    np.add.at(diff, np.arange(len(errors)), errors)
    np.add.at(diff, np.arange(len(errors)) + window, -errors)
    sums = np.cumsum(diff)[:n]
    t = np.arange(n)
    cover = np.minimum(t, n - window) - np.maximum(0, t - window + 1) + 1
    return sums / cover


def pointwise_scores(model, series, device, window=WINDOW):
    errs = window_errors(model, create_windows(series, window, 1), device)
    return pointwise_from_window_errors(errs, len(series), window), errs


def label_free_threshold(support_errors):
    """tau = mean + 2 std of the support-window errors when K > 1, 1.20 x mean when K = 1.
    std is the sample standard deviation (n - 1), as torch's .std() in the original code."""
    e = np.asarray(support_errors, dtype=np.float64)
    return float(e.mean() * 1.20) if len(e) == 1 else float(e.mean() + 2.0 * e.std(ddof=1))


def detection_metrics(scores, labels, tau=None):
    """All metrics for one scored series. F1 at tau is the deployable number; best_f1 uses
    the labels to pick the threshold and is therefore labelled 'oracle'.
    Unlike the original compute_metrics, an all-positive prediction is scored as it is
    (the original set F1 to 0, which hid how much F1 depends on the anomaly fraction)."""
    labels = np.asarray(labels).astype(int)
    scores = np.asarray(scores, dtype=np.float64)
    out = {"n": int(len(labels)), "prevalence": float(labels.mean()),
           "trivial_all_positive_f1": float(2 * labels.mean() / (1 + labels.mean()))}
    both = len(np.unique(labels)) > 1
    out["roc_auc"] = float(roc_auc_score(labels, scores)) if both else None
    out["pr_auc"] = float(average_precision_score(labels, scores)) if both else None
    if both:
        p, r, _ = precision_recall_curve(labels, scores)
        f = 2 * p * r / np.maximum(p + r, 1e-12)
        i = int(np.nanargmax(f))
        out.update(oracle_best_f1=float(f[i]), oracle_best_f1_precision=float(p[i]),
                   oracle_best_f1_recall=float(r[i]))
        out["separation_ratio"] = float(scores[labels == 1].mean() / max(scores[labels == 0].mean(), 1e-12))
    if tau is not None:
        pred = (scores > tau).astype(int)
        out.update(tau=float(tau), flagged_fraction=float(pred.mean()),
                   f1_at_tau=float(f1_score(labels, pred, zero_division=0)),
                   precision_at_tau=float(precision_score(labels, pred, zero_division=0)),
                   recall_at_tau=float(recall_score(labels, pred, zero_division=0)))
    return out


# ---------------------------------------------------------------------------
# Legacy (window-level) query, kept only as a bridge to the old numbers
# ---------------------------------------------------------------------------
def legacy_eval_query(normal_windows, anomaly_windows, k_shot, support_seed, max_normal_query=100):
    """Exactly the original build_eval_query: K normal support windows from the TRAINING
    file; query = ALL anomaly windows from the TEST file + up to 100 other normal windows
    from the training file. Also returns the support/query window indices so overlap
    between support and query windows can be measured."""
    rng = np.random.RandomState(support_seed)
    n_idx = np.arange(len(normal_windows))
    a_idx = np.arange(len(anomaly_windows))
    rng.shuffle(n_idx)
    rng.shuffle(a_idx)
    sup_i = n_idx[:k_shot]
    rem_i = n_idx[k_shot:k_shot + max_normal_query]
    q = np.concatenate([anomaly_windows[a_idx], normal_windows[rem_i]], 0)
    y = np.concatenate([np.ones(len(a_idx)), np.zeros(len(rem_i))])
    perm = rng.permutation(len(q))
    return normal_windows[sup_i], q[perm], y[perm], sup_i, rem_i


def overlapping_fraction(sup_idx, query_idx, starts, window=WINDOW):
    """Fraction of normal query windows that share at least one timestep with a support window."""
    if len(query_idx) == 0 or len(sup_idx) == 0:
        return 0.0
    s = starts[np.asarray(sup_idx)][None, :]
    q = starts[np.asarray(query_idx)][:, None]
    return float((np.abs(q - s) < window).any(axis=1).mean())


# ===========================================================================
# SWaT and WADI
# ===========================================================================
DOWNSAMPLE = 10
PLANT_STRIDE = 10
# Counts found by the earlier work (notebook 02b). The cleaning asserts against these.
SWAT_EXPECTED = {"normal_rows": 495000, "attack_rows": 449919, "attack_labelled_rows": 54621,
                 "sensors": 51}
WADI_EXPECTED = {"normal_rows": 784571, "attack_rows": 172801, "attack_labelled_rows": 9977,
                 "sensors": 123}
WADI_EMPTY_COLS = ["2_LS_001_AL", "2_LS_002_AL", "2_P_001_STATUS", "2_P_002_STATUS"]
WADI_INTERP_COLS = ["1_AIT_002_PV", "1_AIT_004_PV", "2B_AIT_004_PV", "3_AIT_004_PV"]
# Window counts and anomalous-window counts reported by the earlier pipeline.
OLD_WINDOW_COUNTS = {"swat": {"normal": 4948, "attack": 4497, "anomalous": 648},
                     "wadi": {"normal": 7843, "attack": 1726, "anomalous": 140}}
# The regime split used by the original work (swat_task_splits.json, wadi_task_splits.json),
# with the regime sizes it recorded. Notebook_03 rebuilds the regimes, checks that the
# sizes match these, and then uses this split so results stay comparable.
ORIGINAL_REGIMES = {
    "swat": {"k": 12, "meta_train": [1, 9, 5, 2, 10], "meta_val": [0], "meta_test": [4, 8],
             "sizes": {0: 396, 1: 2567, 2: 387, 4: 319, 5: 258, 8: 326, 9: 419, 10: 256}},
    "wadi": {"k": 8, "meta_train": [0, 4, 2, 5, 3], "meta_val": [6], "meta_test": [1, 7],
             "sizes": {0: 133, 1: 958, 2: 1411, 3: 1925, 4: 623, 5: 971, 6: 1491, 7: 331}},
}
PLANT_FILES = {"swat": ("SWaT_Dataset_Normal_v1.xlsx", "SWaT_Dataset_Attack_v0.xlsx"),
               "wadi": ("WADI_14days_new.csv", "WADI_attackdataLABLE.csv")}


def load_swat_workbook(path):
    """Row 1 of the SWaT workbooks is a stage header (P1, P2 ...); row 2 is the real header."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = wb[wb.sheetnames[0]].iter_rows(min_row=2, values_only=True)
    header = [str(h).strip() if h is not None else h for h in next(rows)]
    data = list(rows)
    wb.close()
    label_idx = next(i for i, h in enumerate(header) if h and h.replace(" ", "").lower() == "normal/attack")
    sensor_idx = [i for i in range(len(header)) if i not in (0, label_idx)]
    X = np.array([[r[i] for i in sensor_idx] for r in data], dtype=np.float32)
    labels = [r[label_idx] for r in data]
    return X, labels, [header[i] for i in sensor_idx]


def swat_attack_mask(labels_raw):
    """The attack file spells the label three ways: 'Normal', 'Attack' and 'A ttack'
    (an internal space). All whitespace is removed before matching, and every row must
    then read 'normal' or 'attack'. Returns the mask and the count of each raw spelling."""
    raw = pd.Series([str(v) for v in labels_raw])
    variants = raw.value_counts().to_dict()
    norm = raw.str.replace(r"\s+", "", regex=True).str.lower()
    unknown = set(norm) - {"normal", "attack"}
    assert not unknown, f"unexpected SWaT label values: {unknown}"
    mask = (norm == "attack").to_numpy()
    n_non_normal = sum(c for v, c in variants.items() if v.replace(" ", "").lower() != "normal")
    assert mask.sum() == n_non_normal, "attack count does not match the non-normal label variants"
    return mask, variants


def load_wadi(normal_path, attack_path):
    """WADI.A2: drop the entirely empty columns, interpolate the scattered gaps, and remap
    the attack label (1 = no attack, -1 = attack)."""
    n = pd.read_csv(normal_path, low_memory=False)
    a = pd.read_csv(attack_path, header=1, low_memory=False)     # a stray index row comes first
    n.columns = [c.strip() for c in n.columns]
    a.columns = [c.strip() for c in a.columns]
    a = a.dropna(subset=["Row"])                                 # two blank rows at the end
    label_col = next(c for c in a.columns if "Attack LABLE" in c)
    lab = a[label_col].to_numpy()
    assert set(np.unique(lab)) <= {1, -1}, f"unexpected WADI label values {np.unique(lab)}"
    attack = lab == -1
    meta = ["Row", "Date", "Time"]
    sensors_n = [c for c in n.columns if c not in meta]
    empty = [c for c in sensors_n if n[c].isna().all()]
    assert sorted(empty) == sorted(WADI_EMPTY_COLS), f"empty columns differ: {empty}"
    keep = [c for c in sensors_n if c not in empty]
    assert keep == [c for c in a.columns if c not in meta + [label_col] + empty], "column order differs"
    report = {"empty_columns_dropped": empty, "label_column": label_col}
    out = []
    for name, df in [("normal", n), ("attack", a)]:
        df = df[keep].apply(pd.to_numeric, errors="coerce")
        gappy = [c for c in keep if df[c].isna().any()]
        report[f"{name}_interpolated_columns"] = gappy
        report[f"{name}_missing_values_before"] = int(df.isna().sum().sum())
        for c in gappy:
            df[c] = df[c].interpolate(method="linear", limit_direction="both")
        assert df.isna().sum().sum() == 0
        out.append(df.to_numpy(dtype=np.float32))
    return out[0], out[1], attack, keep, report


def downsample_values(x, factor=DOWNSAMPLE):
    """Every `factor`-th row, after trimming to a multiple of `factor` (as in the earlier work)."""
    n = len(x) // factor * factor
    return x[:n:factor]


def downsample_labels_or(y, factor=DOWNSAMPLE):
    """A downsampled step is an attack if any of its `factor` raw rows is an attack."""
    n = len(y) // factor
    return np.asarray(y[:n * factor]).reshape(n, factor).any(axis=1)


def minmax_scale_plant(Xn, Xa):
    """MinMax fitted on the normal stream only; attack data transformed and clipped to [0, 1]."""
    sc_ = MinMaxScaler().fit(Xn)
    return (sc_.transform(Xn).astype(np.float32),
            np.clip(sc_.transform(Xa), 0, 1).astype(np.float32),
            {"data_min": sc_.data_min_.astype(np.float32), "data_max": sc_.data_max_.astype(np.float32)})


def windows_with_labels(X, y=None, window=WINDOW, stride=PLANT_STRIDE):
    """Windows plus two labels: 'any' (the earlier rule: at least one attack step) and
    'half' (stricter: at least 50% attack steps). Also returns each window's start index."""
    W = create_windows(X, window, stride)
    starts = np.arange(len(W)) * stride
    if y is None:
        return W, starts
    frac = np.array([np.asarray(y[s:s + window]).mean() for s in starts])
    return W, starts, (frac > 0).astype(int), (frac >= 0.5).astype(int)


def regime_summaries(windows):
    """Each window summarised by the per-feature mean and standard deviation over time."""
    return np.concatenate([windows.mean(axis=1), windows.std(axis=1)], axis=1)


def fit_regimes(normal_windows, k_values, seed=42, n_init=10):
    """k-means on standardised window summaries for every k; silhouette per k."""
    from sklearn.cluster import KMeans
    from sklearn.metrics import silhouette_score
    from sklearn.preprocessing import StandardScaler
    S = regime_summaries(normal_windows)
    ss = StandardScaler().fit(S)
    Z = ss.transform(S)
    fits = {}
    for k in k_values:
        km = KMeans(n_clusters=k, n_init=n_init, random_state=seed).fit(Z)
        fits[k] = {"silhouette": float(silhouette_score(Z, km.labels_)), "labels": km.labels_,
                   "centroids": km.cluster_centers_}
    return {"scaler_mean": ss.mean_, "scaler_scale": ss.scale_, "fits": fits,
            "summary_dims": S.shape[1]}


def assign_regimes(windows, scaler_mean, scaler_scale, centroids):
    """Nearest centroid in the same standardised summary space."""
    Z = (regime_summaries(windows) - scaler_mean) / scaler_scale
    d = ((Z[:, None, :] - centroids[None, :, :]) ** 2).sum(-1)
    return d.argmin(1)


def split_regimes(labels, min_size=80, n_val=1, n_test=2, seed=42):
    """Regimes with at least `min_size` normal windows are kept and shuffled with a fixed
    seed; the last `n_test` become meta-test, the `n_val` before them meta-validation."""
    ids, counts = np.unique(labels, return_counts=True)
    sizes = {int(i): int(c) for i, c in zip(ids, counts)}
    kept = sorted(i for i, c in sizes.items() if c >= min_size)
    order = list(np.random.RandomState(seed).permutation(kept))
    return {"meta_train": [int(i) for i in order[:len(order) - n_val - n_test]],
            "meta_val": [int(i) for i in order[len(order) - n_val - n_test:len(order) - n_test]],
            "meta_test": [int(i) for i in order[len(order) - n_test:]],
            "sizes": sizes, "dropped_small": [i for i in sizes if i not in kept], "min_size": min_size}


# ---------------------------------------------------------------------------
# Common low-dimensional space for cross-plant transfer
# ---------------------------------------------------------------------------
def fit_projection(rows_unscaled, n_components=32, seed=42):
    """MinMax scaling, then PCA to `n_components`, then min-max of the projection, all
    fitted on `rows_unscaled` (timesteps x sensors). For the source plant these rows are
    its full normal data; for the target plant, in the leak-free setting, they are only
    the K support windows' rows."""
    from sklearn.decomposition import PCA
    mm = MinMaxScaler().fit(rows_unscaled)
    pca = PCA(n_components=n_components, random_state=seed).fit(mm.transform(rows_unscaled))
    P = pca.transform(mm.transform(rows_unscaled))
    lo, hi = P.min(0), P.max(0)
    return {"minmax": mm, "pca": pca, "lo": lo, "range": np.where(hi - lo > 1e-8, hi - lo, 1.0),
            "explained_variance": float(pca.explained_variance_ratio_.sum()),
            "explained_variance_ratio": pca.explained_variance_ratio_, "n_rows": len(rows_unscaled)}


def apply_projection(proj, rows_unscaled):
    P = proj["pca"].transform(proj["minmax"].transform(rows_unscaled))
    return np.clip((P - proj["lo"]) / proj["range"], 0, 1).astype(np.float32)


def project_windows(proj, windows_unscaled):
    n, T, F = windows_unscaled.shape
    return apply_projection(proj, windows_unscaled.reshape(n * T, F)).reshape(n, T, -1)


def train_on_support(model, support, device, *, seed, steps=300, lr=1e-3):
    """'Scratch': a fresh model trained only on the K support windows, full batch, Adam,
    a fixed number of steps (no validation data exists in this setting)."""
    seed_everything(seed)
    s = to_tensor(support, device)
    opt = torch.optim.Adam(model.parameters(), lr=lr)
    model.train()
    for _ in range(steps):
        opt.zero_grad(); _mse(model(s), s).backward(); opt.step()
    return model


# ---------------------------------------------------------------------------
# Statistics across training seeds
# ---------------------------------------------------------------------------
def describe_values(xs):
    """Mean, sample SD and t-based 95% confidence interval of per-seed values."""
    from scipy import stats as st
    xs = np.array([x for x in xs if x is not None], dtype=float)
    out = {"n": int(len(xs)), "mean": float(xs.mean()) if len(xs) else None}
    if len(xs) > 1:
        sd = xs.std(ddof=1)
        h = st.t.ppf(0.975, len(xs) - 1) * sd / np.sqrt(len(xs))
        out.update(sd=float(sd), ci95=[float(xs.mean() - h), float(xs.mean() + h)])
    return out


def paired_comparison(a, b, margin=0.02):
    """Paired differences a - b over training seeds: mean, 95% CI, two-sided t-test and
    Wilcoxon p-values, and a TOST equivalence p-value for the band [-margin, +margin]
    (p < 0.05 means the difference is shown to lie inside the band)."""
    from scipy import stats as st
    d = np.array([x - y for x, y in zip(a, b) if x is not None and y is not None], dtype=float)
    out = describe_values(d)
    out["margin"] = margin
    if len(d) > 1:
        out["t_test_p_two_sided"] = float(st.ttest_1samp(d, 0.0).pvalue)
        try:
            out["wilcoxon_p_two_sided"] = float(st.wilcoxon(d).pvalue)
        except ValueError:
            out["wilcoxon_p_two_sided"] = None
        se = d.std(ddof=1) / np.sqrt(len(d))
        if se > 0:
            p_low = 1 - st.t.cdf((d.mean() + margin) / se, len(d) - 1)
            p_high = st.t.cdf((d.mean() - margin) / se, len(d) - 1)
            out["tost_p"] = float(max(p_low, p_high))
        out["ci_excludes_zero"] = bool(out["ci95"][0] > 0 or out["ci95"][1] < 0)
    return out


def load_plant(plant, extra_roots=()):
    """Everything Notebook_03 saved for one plant, turned back into windows and regimes."""
    roots = [os.path.join(os.path.dirname(os.path.abspath(__file__)), "outputs"),
             *data_roots(extra_roots)]
    f_clean = find_file(f"{plant}_clean.npz", roots)
    f_reg = find_file(f"{plant}_regimes.npz", roots)
    f_sum = find_file("plant_data_summary.json", roots)
    if not (f_clean and f_reg and f_sum):
        raise FileNotFoundError(f"run Notebook_03 first (or attach its output): {plant}_clean.npz, "
                                f"{plant}_regimes.npz and plant_data_summary.json are needed")
    d = dict(np.load(f_clean, allow_pickle=False))
    r = dict(np.load(f_reg, allow_pickle=False))
    split = json.load(open(f_sum))[plant]["regimes"]["split"]
    Wn, n_starts = windows_with_labels(d["normal_scaled"])
    Wa, a_starts, y_any, y_half = windows_with_labels(d["attack_scaled"], d["attack_labels"])
    Wn_u, _ = windows_with_labels(d["normal_unscaled"])
    Wa_u, _ = windows_with_labels(d["attack_unscaled"])
    assert len(Wn) == len(r["normal_regime"]) and len(Wa) == len(r["attack_regime"])
    return {"plant": plant, "normal": Wn, "attack": Wa, "any": y_any, "half": y_half,
            "normal_unscaled": Wn_u, "attack_unscaled": Wa_u,
            "normal_rows_unscaled": d["normal_unscaled"],
            "normal_regime": r["normal_regime"], "attack_regime": r["attack_regime"],
            "split": {k: split[k] for k in ["meta_train", "meta_val", "meta_test"]},
            "files": [f_clean, f_reg, f_sum]}
